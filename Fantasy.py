
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import tkinter as tk
from tkinter import font as tkfont
from urllib.request import Request, urlopen
from dataclasses import dataclass, replace
from typing import Callable, Iterable
from tkinter import messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


POSITIONS = ("QB", "RB", "WR", "TE")
SLEEPER_PLAYERS_URL = "https://api.sleeper.app/v1/players/nfl"
DEFAULT_EXCEL_FILE = "players.xlsx"


@dataclass(frozen=True)
class Player:
	name: str
	position: str
	global_rank: int
	positional_rank: int
	player_id: str = ""
	team: str = ""
	custom_trade_value: float | None = None


def trade_value(player: Player) -> float:
	if player.global_rank > 150 or player.custom_trade_value is None:
		return 0
	if not 1 <= player.custom_trade_value <= 100:
		return 0
	return player.custom_trade_value


def calculated_trade_value(player: Player) -> float:
	value = trade_value(player)
	if player.global_rank <= 5:
		return value * 2
	if player.global_rank <= 10:
		return value * 1.25
	return value


def trade_verdict(difference: float) -> str:
	if difference <= -30:
		return "They're robbing you"
	if difference <= -15:
		return "Bad trade for you"
	if difference <= -5:
		return "You're giving up too much"
	if difference < 5:
		return "Fair trade"
	if difference < 15:
		return "Good trade for you"
	if difference < 30:
		return "Great trade for you"
	return "You're robbing them"


def find_player(name: str, players: tuple[Player, ...]) -> Player:
	player_lookup = {player.name.casefold(): player for player in players}
	player = player_lookup.get(name.strip().casefold())
	if player is None:
		suggestions = [
			candidate.name
			for candidate in players
			if name.casefold() in candidate.name.casefold()
		]
		hint = f" Did you mean: {', '.join(suggestions)}?" if suggestions else ""
		raise ValueError(f"Player not found: {name}.{hint}")
	return player


def print_rankings(players: tuple[Player, ...], position: str | None, limit: int, sort_by: str) -> None:
	players = [player for player in players if not position or player.position == position]
	if sort_by == "position":
		players.sort(key=lambda player: (player.position, player.positional_rank))
	else:
		players.sort(key=lambda player: player.global_rank)

	print(f"{'Global':<8}{'Pos':<6}{'Pos Rank':<10}{'Player':<24}{'Value':>7}")
	print("-" * 55)
	for player in players[:limit]:
		print(
			f"{player.global_rank:<8}{player.position:<6}"
			f"{player.positional_rank:<10}{player.name:<24}"
			f"{trade_value(player):>7.1f}"
		)


def parse_player_names(names: Iterable[str], players: tuple[Player, ...]) -> list[Player]:
	return [find_player(name, players) for name in names]


def evaluate_trade(give_names: list[str], get_names: list[str], players: tuple[Player, ...]) -> None:
	giving = parse_player_names(give_names, players)
	receiving = parse_player_names(get_names, players)
	trade_players = (*giving, *receiving)
	if any(player.global_rank > 150 for player in trade_players):
		raise ValueError("Trades are limited to the top 150 global players.")
	if any(trade_value(player) == 0 for player in trade_players):
		raise ValueError("Set each player's trade value from 1 to 100 in Edit Rankings before calculating.")
	give_total = sum(trade_value(player) for player in giving)
	get_total = sum(trade_value(player) for player in receiving)
	difference = get_total - give_total

	print("You give:")
	for player in giving:
		print(f"  {player.name} ({player.position}{player.positional_rank}) = {trade_value(player):.1f}")
	print("You get:")
	for player in receiving:
		print(f"  {player.name} ({player.position}{player.positional_rank}) = {trade_value(player):.1f}")
	print(f"\nTotals: {give_total:.1f} given | {get_total:.1f} received")
	print(f"Difference: {difference:+.1f} ({trade_verdict(difference)})")


def fetch_api_players() -> tuple[Player, ...]:
	request = Request(SLEEPER_PLAYERS_URL, headers={"User-Agent": "FantasyPy/1.0"})
	with urlopen(request, timeout=30) as response:
		data = json.load(response)

	api_players = []
	for player in data.values():
		position = player.get("position")
		search_rank = player.get("search_rank")
		if position not in POSITIONS or not player.get("active") or search_rank is None:
			continue
		name = " ".join(filter(None, (player.get("first_name"), player.get("last_name"))))
		if name:
			api_players.append((name, position, int(search_rank), player["player_id"], player.get("team") or "FA"))

	api_players.sort(key=lambda player: player[2])
	position_counts = {position: 0 for position in POSITIONS}
	players = []
	for global_rank, (name, position, _, player_id, team) in enumerate(api_players, start=1):
		position_counts[position] += 1
		players.append(Player(name, position, global_rank, position_counts[position], player_id, team))
	return tuple(players)


def write_excel(players: tuple[Player, ...], output_path: str) -> None:
	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "Players"
	worksheet.append(["Name", "Position", "Team", "Global Rank", "Positional Rank", "Trade Value", "Player ID", "Custom Trade Value"])
	for player in players:
		worksheet.append([
			player.name,
			player.position,
			player.team,
			player.global_rank,
			player.positional_rank,
			trade_value(player),
			player.player_id,
			player.custom_trade_value,
		])
	for cell in worksheet[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="1F4E78")
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	for column, width in {"A": 24, "B": 12, "C": 10, "D": 14, "E": 18, "F": 14, "G": 14, "H": 20}.items():
		worksheet.column_dimensions[column].width = width
	workbook.save(output_path)


def read_excel(input_path: str) -> tuple[Player, ...]:
	workbook = load_workbook(input_path, read_only=True, data_only=True)
	worksheet = workbook["Players"]
	columns = {cell.value: cell.column - 1 for cell in worksheet[1]}
	players = []
	for row in worksheet.iter_rows(min_row=2, values_only=True):
		if not row[columns["Name"]]:
			continue
		custom_value_index = columns.get("Custom Trade Value")
		custom_value = row[custom_value_index] if custom_value_index is not None else None
		players.append(Player(
			str(row[columns["Name"]]),
			str(row[columns["Position"]]),
			int(row[columns["Global Rank"]]),
			int(row[columns["Positional Rank"]]),
			str(row[columns["Player ID"]] or ""),
			str(row[columns["Team"]] or ""),
			float(custom_value) if custom_value not in (None, "") else None,
		))
	workbook.close()
	return tuple(players)


def save_rankings(input_path: str, players: tuple[Player, ...]) -> None:
	workbook = load_workbook(input_path)
	worksheet = workbook["Players"]
	columns = {cell.value: cell.column for cell in worksheet[1]}
	global_column = columns["Global Rank"]
	positional_column = columns["Positional Rank"]
	value_column = columns["Trade Value"]
	custom_value_column = columns.get("Custom Trade Value")
	if custom_value_column is None:
		custom_value_column = worksheet.max_column + 1
		header = worksheet.cell(1, custom_value_column, "Custom Trade Value")
		header.font = Font(bold=True, color="FFFFFF")
		header.fill = PatternFill("solid", fgColor="1F4E78")
		worksheet.column_dimensions[header.column_letter].width = 20
	id_column = columns["Player ID"]
	name_column = columns["Name"]
	position_column = columns["Position"]
	team_column = columns["Team"]
	player_lookup = {
		("id", player.player_id) if player.player_id else ("name", player.name.casefold(), player.position, player.team.casefold()): player
		for player in players
	}
	updated = 0
	for row in range(2, worksheet.max_row + 1):
		player_id = worksheet.cell(row, id_column).value
		if player_id is not None:
			key = ("id", str(player_id))
		else:
			name = worksheet.cell(row, name_column).value
			position = worksheet.cell(row, position_column).value
			team = worksheet.cell(row, team_column).value or ""
			key = ("name", str(name).casefold(), str(position), str(team).casefold())
		player = player_lookup.get(key)
		if player is None:
			continue
		worksheet.cell(row, global_column, player.global_rank)
		worksheet.cell(row, positional_column, player.positional_rank)
		worksheet.cell(row, value_column, trade_value(player))
		worksheet.cell(row, custom_value_column, player.custom_trade_value)
		updated += 1
	if updated != len(player_lookup):
		workbook.close()
		raise ValueError("Could not match every ranked player to a row in the workbook.")
	workbook.save(input_path)
	workbook.close()


def load_players(input_path: str | None) -> tuple[Player, ...]:
	return read_excel(input_path or DEFAULT_EXCEL_FILE)



class RoundedButton(tk.Canvas):
	def __init__(self, parent: tk.Misc, text: str, command: Callable[[], None], width: int = 170, height: int = 40) -> None:
		super().__init__(parent, width=width, height=height, background="#151515", highlightthickness=0, cursor="hand2")
		self.text = text
		self.command = command
		self.width = width
		self.height = height
		self.selected = False
		self.hovered = False
		self.bind("<Button-1>", self._activate)
		self.bind("<Enter>", self._on_enter)
		self.bind("<Leave>", self._on_leave)
		self._draw()

	def set_selected(self, selected: bool) -> None:
		self.selected = selected
		self._draw()

	def _activate(self, _event: tk.Event) -> None:
		self.command()

	def _on_enter(self, _event: tk.Event) -> None:
		self.hovered = True
		self._draw()

	def _on_leave(self, _event: tk.Event) -> None:
		self.hovered = False
		self._draw()

	def _draw(self) -> None:
		self.delete("all")
		fill = "#F47721" if self.selected else "#383838" if self.hovered else "#202020"
		foreground = "#111111" if self.selected else "#F4F4F4"
		radius = min(10, self.height // 2)
		self.create_polygon(
			radius, 0, self.width - radius, 0, self.width, radius,
			self.width, self.height - radius, self.width - radius, self.height,
			radius, self.height, 0, self.height - radius, 0, radius,
			fill=fill,
			outline="",
			smooth=True,
			splinesteps=16,
		)
		self.create_text(self.width / 2, self.height / 2, text=self.text, fill=foreground, font=("Segoe UI", 10, "bold"))


class RoundedField(tk.Canvas):
	def __init__(self, parent: tk.Misc, textvariable: tk.StringVar, width: int, height: int = 34) -> None:
		super().__init__(parent, width=width, height=height, background="#151515", highlightthickness=0)
		self.create_polygon(
			10, 0, width - 10, 0, width, 10, width, height - 10,
			width - 10, height, 10, height, 0, height - 10, 0, 10,
			fill="#202020",
			outline="#383838",
			smooth=True,
			splinesteps=16,
		)
		self.entry = ttk.Entry(self, textvariable=textvariable, width=18, style="Rounded.TEntry")
		self.create_window(5, 3, window=self.entry, width=width - 10, height=height - 6, anchor="nw")
		self.bind("<Button-1>", lambda _event: self.entry.focus_set())


class TradePlayerBubble(tk.Canvas):
	def __init__(self, parent: tk.Misc, player: Player, remove_player: Callable[[Player], None]) -> None:
		self.name_font = tkfont.Font(parent, family="Segoe UI", size=10, weight="bold")
		self.metadata_font = tkfont.Font(parent, family="Segoe UI", size=9)
		metadata = f"{player.team}  |  {player.position}"
		width = max(150, min(320, max(self.name_font.measure(player.name), self.metadata_font.measure(metadata)) + 48))
		super().__init__(parent, width=width, height=58, background="#151515", highlightthickness=0)
		self.width = width
		self.create_polygon(
			14, 0, width - 14, 0, width, 14, width, 44,
			width - 14, 58, 14, 58, 0, 44, 0, 14,
			fill="#242424",
			outline="#383838",
			smooth=True,
			splinesteps=16,
		)
		self.create_text(12, 19, text=player.name, fill="#F4F4F4", font=self.name_font, anchor="w")
		self.create_text(12, 41, text=metadata, fill="#B8B8B8", font=self.metadata_font, anchor="w")
		self.create_text(width - 17, 29, text="x", fill="#F47721", font=("Segoe UI", 10, "bold"), tags=("remove",))
		self.tag_bind("remove", "<Button-1>", lambda _event: remove_player(player))
		self.tag_bind("remove", "<Enter>", lambda _event: self.configure(cursor="hand2"))
		self.tag_bind("remove", "<Leave>", lambda _event: self.configure(cursor=""))


class TradeBubbleWrap(ttk.Frame):
	def __init__(self, parent: tk.Misc) -> None:
		super().__init__(parent)
		self.bubbles: list[TradePlayerBubble] = []
		self.bind("<Configure>", self._layout_bubbles)

	def set_players(self, players: list[Player], remove_player: Callable[[Player], None]) -> None:
		for bubble in self.bubbles:
			bubble.destroy()
		self.bubbles = [TradePlayerBubble(self, player, remove_player) for player in players]
		self.after_idle(self._layout_bubbles)

	def _layout_bubbles(self, _event: tk.Event | None = None) -> None:
		available_width = self.winfo_width()
		if available_width <= 1:
			return
		x = 0
		y = 0
		for bubble in self.bubbles:
			if x and x + bubble.width > available_width:
				x = 0
				y += 66
			bubble.place(x=x, y=y, width=bubble.width, height=58)
			x += bubble.width + 8


class RankingsWindow:
	def __init__(self, root: tk.Tk, excel_path: str) -> None:
		self.root = root
		self.excel_path = excel_path
		self.players: tuple[Player, ...] = ()
		self.dirty = False
		self.selected_player_index: int | None = None
		self.search = tk.StringVar()
		self.edit_global_rank = tk.StringVar()
		self.edit_positional_rank = tk.StringVar()
		self.edit_trade_value = tk.StringVar()
		self.selected_player_name = tk.StringVar(value="Select a player from the list")
		self.trade_sides: dict[str, list[Player]] = {"give": [], "get": []}
		self.trade_entries: dict[str, ttk.Entry] = {}
		self.trade_suggestions: dict[str, ttk.Treeview] = {}
		self.trade_matches: dict[str, list[Player]] = {}
		self.trade_bubbles: dict[str, TradeBubbleWrap] = {}
		root.title("FantasyPy Rankings")
		root.geometry("1000x720")
		root.minsize(760, 660)
		self.build_layout()
		self.load_rankings()

	def build_layout(self) -> None:
		style = ttk.Style()
		style.theme_use("clam")
		background = "#151515"
		panel = "#202020"
		orange = "#F47721"
		white = "#F4F4F4"
		muted = "#B8B8B8"
		self.root.configure(background=background)
		style.configure("TFrame", background=background)
		style.configure("TLabel", background=background, foreground=white, font=("Segoe UI", 10))
		style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=(12, 7), background=panel, foreground=white)
		style.map("TButton", background=[("active", "#383838")], foreground=[("active", white)])
		style.configure("Accent.TButton", background=orange, foreground="#111111")
		style.map("Accent.TButton", background=[("active", "#FF964C")], foreground=[("active", "#111111")])
		style.configure("TEntry", fieldbackground=panel, foreground=white, insertcolor=white)
		style.configure("Rounded.TEntry", fieldbackground=panel, foreground=white, insertcolor=white, borderwidth=0, padding=(4, 2))
		style.configure("TCombobox", fieldbackground=panel, background=panel, foreground=white, arrowcolor=orange)
		style.map("TCombobox", fieldbackground=[("readonly", panel)], foreground=[("readonly", white)])
		style.configure("Treeview", rowheight=30, font=("Segoe UI", 10), background=panel, fieldbackground=panel, foreground=white)
		style.map("Treeview", background=[("selected", orange)], foreground=[("selected", "#111111")])
		style.configure("Compact.Treeview", rowheight=18, font=("Segoe UI", 9), background=panel, fieldbackground=panel, foreground=white)
		style.map("Compact.Treeview", background=[("selected", orange)], foreground=[("selected", "#111111")])
		style.configure("Compact.Treeview.Heading", font=("Segoe UI", 8, "bold"), background="#101010", foreground=orange, padding=(5, 3))
		style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"), background="#101010", foreground=orange, padding=(8, 9))
		style.map("Treeview.Heading", background=[("active", "#303030")])
		header = ttk.Frame(self.root, padding=(18, 16, 18, 8))
		header.pack(fill="x")
		tk.Label(header, text="FantasyPy Rankings", background=background, foreground=white, font=("Segoe UI", 21, "bold")).pack(anchor="w")
		tk.Label(header, text=f"SOURCE  /  {Path(self.excel_path).name}", background=background, foreground=muted).pack(anchor="w", pady=(3, 0))

		navigation = ttk.Frame(self.root, padding=(18, 8, 18, 0))
		navigation.pack(fill="x")
		self.content = ttk.Frame(self.root)
		self.content.pack(fill="both", expand=True, padx=18, pady=(8, 0))
		self.content.rowconfigure(0, weight=1)
		self.content.columnconfigure(0, weight=1)
		self.trade_tab = ttk.Frame(self.content)
		self.global_tab = ttk.Frame(self.content)
		self.positional_tab = ttk.Frame(self.content)
		self.edit_tab = ttk.Frame(self.content)
		self.navigation_buttons = []
		self.navigation_pages = []
		for label, page in (
			("Trade Calculator", self.trade_tab),
			("Global Rankings", self.global_tab),
			("Positional Rankings", self.positional_tab),
			("Edit Rankings", self.edit_tab),
		):
			page.grid(row=0, column=0, sticky="nsew")
			button = RoundedButton(navigation, label, lambda selected_page=page: self.show_page(selected_page))
			button.pack(side="left", padx=(0, 8))
			self.navigation_buttons.append(button)
			self.navigation_pages.append(page)
		self.show_page(self.trade_tab)
		self.build_trade_calculator()

		global_controls = ttk.Frame(self.global_tab, padding=(0, 8))
		global_controls.pack(fill="x")
		tk.Label(global_controls, text="Search", background=background, foreground=white).pack(side="left")
		search_field = RoundedField(global_controls, self.search, width=160)
		search_field.pack(side="left", padx=(6, 8))
		search_field.entry.bind("<KeyRelease>", lambda _event: self.refresh_table())
		RoundedButton(global_controls, text="Refresh", command=self.load_rankings, width=82, height=34).pack(side="right", padx=(8, 0))
		RoundedButton(global_controls, text="Open Excel", command=self.open_excel, width=104, height=34).pack(side="right")

		self.table = self.create_rankings_table(self.global_tab, include_positional_rank=False)
		self.positional_tables: dict[str, ttk.Treeview] = {}
		self.build_positional_board()

		edit_content = ttk.Frame(self.edit_tab, padding=(0, 8, 0, 12))
		edit_content.pack(fill="both", expand=True)
		ttk.Label(edit_content, text="Select a player to edit their global and positional ranks.", foreground=muted).pack(anchor="w", pady=(0, 8))
		self.editor_table = self.create_rankings_table(edit_content, include_trade_value=True)
		self.editor_table.bind("<<TreeviewSelect>>", self.select_player)

		player_form = ttk.Frame(edit_content, padding=(14, 12))
		player_form.pack(fill="x", pady=(12, 0))
		ttk.Label(player_form, textvariable=self.selected_player_name, font=("Segoe UI", 12, "bold")).grid(row=0, column=0, columnspan=5, sticky="w", pady=(0, 10))
		ttk.Label(player_form, text="Global Rank").grid(row=1, column=0, sticky="w")
		global_rank_entry = ttk.Entry(player_form, textvariable=self.edit_global_rank, width=10)
		global_rank_entry.grid(row=2, column=0, sticky="w", padx=(0, 18))
		tk.Label(player_form, text="Positional Rank").grid(row=1, column=1, sticky="w")
		positional_rank_entry = ttk.Entry(player_form, textvariable=self.edit_positional_rank, width=10)
		positional_rank_entry.grid(row=2, column=1, sticky="w", padx=(0, 18))
		tk.Label(player_form, text="Trade Value").grid(row=1, column=2, sticky="w")
		self.trade_value_entry = ttk.Entry(player_form, textvariable=self.edit_trade_value, width=10)
		self.trade_value_entry.grid(row=2, column=2, sticky="w", padx=(0, 18))
		self.edit_trade_value.trace_add("write", lambda *_args: self.preview_trade_value())
		self.apply_button = ttk.Button(player_form, text="Apply to Player", command=self.apply_rank_edits)
		self.apply_button.grid(row=2, column=3, sticky="w", padx=(0, 8))
		ttk.Button(player_form, text="Save Changes", style="Accent.TButton", command=self.save_changes).grid(row=2, column=4, sticky="w")
		player_form.columnconfigure(5, weight=1)

		self.status = ttk.Label(self.root, padding=(18, 0, 18, 10))
		self.status.pack(fill="x")

	def show_page(self, page: ttk.Frame) -> None:
		page.tkraise()
		for button, candidate in zip(self.navigation_buttons, self.navigation_pages):
			button.set_selected(candidate is page)

	def build_trade_calculator(self) -> None:
		board = ttk.Frame(self.trade_tab, padding=(0, 8))
		board.pack(fill="both", expand=True)
		board.columnconfigure(0, weight=1, uniform="trade-side")
		board.columnconfigure(2, weight=1, uniform="trade-side")
		board.rowconfigure(0, weight=1)
		ttk.Separator(board, orient="vertical").grid(row=0, column=1, sticky="ns", padx=10)

		for column, (side, title) in enumerate((("give", "YOU GIVE"), ("get", "YOU GET"))):
			grid_column = column * 2
			panel = ttk.Frame(board, padding=10)
			panel.grid(row=0, column=grid_column, sticky="nsew")
			panel.rowconfigure(3, weight=1)
			panel.columnconfigure(0, weight=1)
			ttk.Label(panel, text=title, foreground="#F47721", font=("Segoe UI", 13, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 8))

			add_row = ttk.Frame(panel)
			add_row.grid(row=1, column=0, sticky="ew", pady=(0, 10))
			add_row.columnconfigure(0, weight=1)
			entry = ttk.Entry(add_row)
			entry.grid(row=0, column=0, sticky="ew")
			entry.bind("<Return>", lambda _event, selected_side=side: self.add_trade_player(selected_side))
			entry.bind("<KeyRelease>", lambda _event, selected_side=side: self.update_trade_suggestions(selected_side))
			self.trade_entries[side] = entry
			ttk.Button(add_row, text="Add Player", style="Accent.TButton", command=lambda selected_side=side: self.add_trade_player(selected_side)).grid(row=0, column=1, padx=(8, 0))

			suggestion_frame = ttk.Frame(panel)
			suggestion_frame.grid(row=2, column=0, sticky="ew", pady=(0, 8))
			suggestion_frame.columnconfigure(0, weight=1)
			suggestions = ttk.Treeview(suggestion_frame, columns=("player", "team", "position"), show="headings", selectmode="browse", height=4, style="Compact.Treeview")
			for column_id, heading, width in (("player", "Player", 150), ("team", "Team", 55), ("position", "Pos", 48)):
				suggestions.heading(column_id, text=heading)
				suggestions.column(column_id, width=width, minwidth=40, stretch=column_id == "player", anchor="w" if column_id == "player" else "center")
			suggestion_scrollbar = ttk.Scrollbar(suggestion_frame, orient="vertical", command=suggestions.yview)
			suggestions.configure(yscrollcommand=suggestion_scrollbar.set)
			suggestions.grid(row=0, column=0, sticky="ew")
			suggestion_scrollbar.grid(row=0, column=1, sticky="ns")
			suggestions.bind("<Double-1>", lambda _event, selected_side=side: self.add_trade_player(selected_side))
			suggestions.bind("<Return>", lambda _event, selected_side=side: self.add_trade_player(selected_side))
			suggestion_frame.grid_remove()
			self.trade_suggestions[side] = suggestions
			self.trade_matches[side] = []

			bubbles = TradeBubbleWrap(panel)
			bubbles.grid(row=3, column=0, sticky="nsew")
			self.trade_bubbles[side] = bubbles

		ttk.Button(board, text="Calculate Trade", style="Accent.TButton", command=self.calculate_trade).grid(row=1, column=0, columnspan=3, pady=(14, 4))

	def add_trade_player(self, side: str) -> None:
		query = self.trade_entries[side].get().strip()
		selection = self.trade_suggestions[side].selection()
		player = self.trade_matches[side][int(selection[0])] if selection else None
		if player is None:
			if not query:
				return
			available_players = tuple(player for player in self.players if player.global_rank <= 150 and player.team.casefold() != "fa")
			player_lookup = {candidate.name.casefold(): candidate for candidate in available_players}
			player = player_lookup.get(query.casefold())
			if player is None:
				matches = [candidate for candidate in available_players if query.casefold() in candidate.name.casefold()]
				if len(matches) == 1:
					player = matches[0]
				elif matches:
					names = ", ".join(candidate.name for candidate in matches[:8])
					more = f" and {len(matches) - 8} more" if len(matches) > 8 else ""
					messagebox.showinfo("Choose a player", f"Several players match: {names}{more}. Enter a more specific name.")
					return
				else:
					messagebox.showerror("Player not found", f"No player matches '{query}'.")
					return

		def is_same_player(candidate: Player) -> bool:
			if player.player_id and candidate.player_id:
				return player.player_id == candidate.player_id
			return (
				player.name.casefold(), player.position, player.team.casefold()
			) == (
				candidate.name.casefold(), candidate.position, candidate.team.casefold()
			)

		if any(is_same_player(candidate) for players in self.trade_sides.values() for candidate in players):
			messagebox.showinfo("Player already added", "A player can only appear once in a trade.")
			return
		self.trade_sides[side].append(player)
		self.trade_entries[side].delete(0, tk.END)
		self.refresh_trade_side(side)
		self.update_trade_suggestions(side)

	def update_trade_suggestions(self, side: str) -> None:
		suggestions = self.trade_suggestions[side]
		for item in suggestions.get_children():
			suggestions.delete(item)
		query = self.trade_entries[side].get().strip().casefold()
		if not query:
			self.trade_matches[side] = []
			suggestions.master.grid_remove()
			return
		selected = {self.trade_player_key(player) for players in self.trade_sides.values() for player in players}
		matches = [
			player for player in self.players
			if player.global_rank <= 150
			and player.team.casefold() != "fa"
			and query in player.name.casefold()
			and self.trade_player_key(player) not in selected
		][:30]
		self.trade_matches[side] = matches
		if not matches:
			suggestions.master.grid_remove()
			return
		suggestions.configure(height=min(4, len(matches)))
		for index, player in enumerate(matches):
			suggestions.insert("", "end", iid=str(index), values=(player.name, player.team, player.position))
		suggestions.master.grid()

	def remove_trade_player(self, side: str, player: Player) -> None:
		if player not in self.trade_sides[side]:
			return
		self.trade_sides[side].remove(player)
		self.refresh_trade_side(side)
		self.update_trade_suggestions(side)

	def refresh_trade_side(self, side: str) -> None:
		self.trade_bubbles[side].set_players(
			self.trade_sides[side],
			lambda player: self.remove_trade_player(side, player),
		)

	def calculate_trade(self) -> None:
		if not self.trade_sides["give"] or not self.trade_sides["get"]:
			messagebox.showinfo("Trade incomplete", "Add at least one player to each side before calculating.")
			return
		players = (*self.trade_sides["give"], *self.trade_sides["get"])
		if any(player.global_rank > 150 for player in players):
			messagebox.showinfo("Top 150 only", "Trades are limited to the top 150 global players.")
			return
		if any(trade_value(player) == 0 for player in players):
			messagebox.showinfo("Trade values needed", "Set each player's value from 1 to 100 in Edit Rankings before calculating.")
			return
		give_total = sum(calculated_trade_value(player) for player in self.trade_sides["give"])
		get_total = sum(calculated_trade_value(player) for player in self.trade_sides["get"])
		difference = get_total - give_total
		self.show_trade_result(difference)

	def show_trade_result(self, difference: float) -> None:
		popup = tk.Toplevel(self.root)
		popup.title("Trade Result")
		popup.configure(background="#151515")
		popup.resizable(False, False)
		popup.transient(self.root)

		content = ttk.Frame(popup, padding=(20, 18))
		content.pack(fill="both", expand=True)
		ttk.Label(content, text=f"Difference: {difference:+.1f}", foreground="#F47721", font=("Segoe UI", 24, "bold")).pack()
		ttk.Label(content, text=trade_verdict(difference), font=("Segoe UI", 13, "bold")).pack(pady=(4, 8))

		scale = tk.Canvas(content, width=460, height=78, background="#151515", highlightthickness=0)
		scale.pack(fill="x", pady=(4, 10))
		left = 10
		right = 450
		scale_limit = max(40, abs(difference))
		colors = ("#A83E3E", "#C35A48", "#CF8C4D", "#777777", "#A2A14E", "#70A65D", "#3C9461")
		bounds = (-scale_limit, -30, -15, -5, 5, 15, 30, scale_limit)
		for start, end, color in zip(bounds, bounds[1:], colors):
			x1 = left + (start + scale_limit) / (2 * scale_limit) * (right - left)
			x2 = left + (end + scale_limit) / (2 * scale_limit) * (right - left)
			scale.create_rectangle(x1, 10, x2, 27, fill=color, outline="#151515", width=1)

		marker_x = left + (difference + scale_limit) / (2 * scale_limit) * (right - left)
		scale.create_line(marker_x, 3, marker_x, 33, fill="#FFFFFF", width=3)
		for tick in (-30, -15, -5, 5, 15, 30):
			tick_x = left + (tick + scale_limit) / (2 * scale_limit) * (right - left)
			scale.create_text(tick_x, 43, text=f"{tick:+}", fill="#B8B8B8", font=("Segoe UI", 8))
		scale.create_text(left, 64, text="Bad for you", fill="#B8B8B8", font=("Segoe UI", 9), anchor="w")
		scale.create_text((left + right) / 2, 64, text="Fair", fill="#B8B8B8", font=("Segoe UI", 9), anchor="center")
		scale.create_text(right, 64, text="Good for you", fill="#B8B8B8", font=("Segoe UI", 9), anchor="e")

		ttk.Button(content, text="Close", command=popup.destroy).pack(anchor="e")
		popup.bind("<Escape>", lambda _event: popup.destroy())
		popup.update_idletasks()
		x = self.root.winfo_rootx() + (self.root.winfo_width() - popup.winfo_width()) // 2
		y = self.root.winfo_rooty() + (self.root.winfo_height() - popup.winfo_height()) // 2
		popup.geometry(f"+{x}+{y}")
		popup.grab_set()
		popup.focus_set()

	def trade_player_key(self, player: Player) -> tuple[str, str, str]:
		return (player.player_id or player.name.casefold(), player.position, player.team.casefold())

	def refresh_trade_selections(self) -> None:
		current_players = {self.trade_player_key(player): player for player in self.players}
		for side in self.trade_sides:
			self.trade_sides[side] = [
				current_players[self.trade_player_key(player)]
				for player in self.trade_sides[side]
				if self.trade_player_key(player) in current_players
			]
			self.refresh_trade_side(side)

	def create_rankings_table(
		self,
		parent: ttk.Frame,
		include_positional_rank: bool = True,
		include_trade_value: bool = False,
	) -> ttk.Treeview:
		table_frame = ttk.Frame(parent)
		table_frame.pack(fill="both", expand=True)
		columns = ["global", "position"]
		column_details = [("global", "Global", 80), ("position", "Position", 90)]
		if include_positional_rank:
			columns.append("positional")
			column_details.append(("positional", "Pos. Rank", 100))
		if include_trade_value:
			columns.append("trade_value")
			column_details.append(("trade_value", "Value", 75))
		columns.extend(("name", "team"))
		column_details.extend((("name", "Player", 280), ("team", "Team", 90)))
		table = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
		for column, heading, width in column_details:
			table.heading(column, text=heading)
			table.column(column, width=width, anchor="center" if column != "name" else "w")
		scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=table.yview)
		table.configure(yscrollcommand=scrollbar.set)
		table.pack(side="left", fill="both", expand=True)
		scrollbar.pack(side="right", fill="y")
		return table

	def build_positional_board(self) -> None:
		board = ttk.Frame(self.positional_tab, padding=(0, 4, 0, 6))
		board.pack(fill="both", expand=True)
		for index, position in enumerate(POSITIONS):
			group = ttk.Frame(board, padding=4)
			group.grid(row=index // 2, column=index % 2, sticky="nsew", padx=5, pady=5)
			ttk.Label(group, text=position, foreground="#F47721", font=("Segoe UI", 13, "bold")).pack(anchor="w", pady=(0, 5))
			table = ttk.Treeview(group, columns=("rank", "name", "team"), show="headings", height=10, style="Compact.Treeview")
			for column, heading, width in (("rank", "Rank", 58), ("name", "Player", 230), ("team", "Team", 70)):
				table.heading(column, text=heading)
				table.column(column, width=width, anchor="center" if column != "name" else "w")
			table.pack(fill="both", expand=True)
			self.positional_tables[position] = table
		for row in range(2):
			board.rowconfigure(row, weight=1)
		for column in range(2):
			board.columnconfigure(column, weight=1)

	def load_rankings(self) -> None:
		if self.dirty and not messagebox.askyesno("Discard edits?", "You have unsaved ranking changes. Discard them and reload the workbook?"):
			return
		try:
			self.players = load_players(self.excel_path)
			self.dirty = False
			self.refresh_trade_selections()
			self.refresh_table()
			self.refresh_positional_tables()
			self.refresh_editor_table()
		except (OSError, KeyError, TypeError, ValueError) as error:
			messagebox.showerror("Could not load rankings", str(error))

	def refresh_table(self) -> None:
		players = [
			(index, player)
			for index, player in enumerate(self.players)
			if player.team.casefold() != "fa"
		]
		query = self.search.get().strip().casefold()
		if query:
			players = [(index, player) for index, player in players if query in player.name.casefold() or query in player.team.casefold()]
		players.sort(key=lambda item: item[1].global_rank)
		for item in self.table.get_children():
			self.table.delete(item)
		for index, player in players:
			self.table.insert("", "end", iid=str(index), values=(player.global_rank, player.position, player.name, player.team))
		self.style_table_rows(self.table)
		message = "Unsaved changes" if self.dirty else ""
		self.status.configure(text=f"Showing {len(players):,} players  |  {message}")

	def refresh_positional_tables(self) -> None:
		for position, table in self.positional_tables.items():
			players = [player for player in self.players if player.position == position and player.team.casefold() != "fa"]
			players.sort(key=lambda player: (player.positional_rank, player.global_rank))
			for item in table.get_children():
				table.delete(item)
			for player in players[:10]:
				table.insert("", "end", values=(player.positional_rank, player.name, player.team))
			self.style_table_rows(table)

	def refresh_editor_table(self, selected_index: int | None = None) -> None:
		for item in self.editor_table.get_children():
			self.editor_table.delete(item)
		players = [(index, player) for index, player in enumerate(self.players) if player.team.casefold() != "fa"]
		players.sort(key=lambda item: item[1].global_rank)
		for index, player in players:
			value = trade_value(player)
			value_display = f"{value:g}" if value else "N/A" if player.global_rank > 150 else "Unset"
			self.editor_table.insert(
				"",
				"end",
				iid=str(index),
				values=(player.global_rank, player.position, player.positional_rank, value_display, player.name, player.team),
			)
		self.style_table_rows(self.editor_table)
		if selected_index is not None and self.editor_table.exists(str(selected_index)):
			self.editor_table.selection_set(str(selected_index))
			self.editor_table.focus(str(selected_index))
			self.set_selected_player(selected_index)
		else:
			self.selected_player_index = None
			self.selected_player_name.set("Select a player from the list")
			self.edit_global_rank.set("")
			self.edit_positional_rank.set("")
			self.edit_trade_value.set("")
			self.trade_value_entry.configure(state="disabled")

	def style_table_rows(self, table: ttk.Treeview) -> None:
		table.tag_configure("even", background="#202020")
		table.tag_configure("odd", background="#292929")
		for index, item in enumerate(table.get_children()):
			table.item(item, tags=("even" if index % 2 == 0 else "odd",))

	def select_player(self, _event: tk.Event) -> None:
		selection = self.editor_table.selection()
		if selection:
			self.set_selected_player(int(selection[0]))

	def set_selected_player(self, player_index: int) -> None:
		self.selected_player_index = player_index
		player = self.players[player_index]
		self.selected_player_name.set(f"{player.name}  /  {player.position}  /  {player.team}")
		self.edit_global_rank.set(str(player.global_rank))
		self.edit_positional_rank.set(str(player.positional_rank))
		value = trade_value(player)
		self.edit_trade_value.set(f"{value:g}" if value else "")
		self.trade_value_entry.configure(state="normal" if player.global_rank <= 150 else "disabled")

	def preview_trade_value(self) -> None:
		if self.selected_player_index is None:
			return
		item = str(self.selected_player_index)
		if not self.editor_table.exists(item):
			return
		try:
			rank = int(self.edit_global_rank.get())
		except ValueError:
			rank = self.players[self.selected_player_index].global_rank
		if rank > 150:
			display = "N/A"
		else:
			text = self.edit_trade_value.get().strip()
			if not text:
				display = "Unset"
			else:
				try:
					value = float(text)
					display = f"{value:g}" if math.isfinite(value) and 1 <= value <= 100 else "Invalid"
				except ValueError:
					display = "Invalid"
		self.editor_table.set(item, "trade_value", display)

	def apply_rank_edits(self) -> None:
		if self.selected_player_index is None:
			messagebox.showinfo("Select a player", "Select a player before applying rank changes.")
			return
		try:
			global_rank = int(self.edit_global_rank.get())
			positional_rank = int(self.edit_positional_rank.get())
			trade_value_text = self.edit_trade_value.get().strip()
			custom_trade_value = float(trade_value_text) if trade_value_text else None
			if (
				global_rank < 1
				or positional_rank < 1
				or (custom_trade_value is not None and (not math.isfinite(custom_trade_value) or not 1 <= custom_trade_value <= 100))
			):
				raise ValueError
		except ValueError:
			messagebox.showerror("Invalid values", "Ranks must be whole numbers above zero. Trade value must be blank or from 1 to 100.")
			return
		players = list(self.players)
		original_player = players[self.selected_player_index]
		players[self.selected_player_index] = replace(
			original_player,
			global_rank=global_rank,
			positional_rank=positional_rank,
			custom_trade_value=custom_trade_value if global_rank <= 150 else None,
		)
		self.players = tuple(players)
		self.dirty = True
		self.refresh_trade_selections()
		self.refresh_table()
		self.refresh_positional_tables()
		self.refresh_editor_table(self.selected_player_index)
		self.status.configure(text="Rank changes applied. Select Save Changes to write them to Excel.")

	def save_changes(self) -> None:
		if not self.dirty:
			self.status.configure(text="No unsaved changes")
			return
		try:
			save_rankings(self.excel_path, self.players)
		except (OSError, KeyError, TypeError, ValueError) as error:
			messagebox.showerror("Could not save rankings", str(error))
			return
		self.dirty = False
		self.refresh_table()
		self.refresh_positional_tables()
		self.refresh_editor_table(self.selected_player_index)
		self.status.configure(text=f"Rankings saved to {Path(self.excel_path).name}")

	def open_excel(self) -> None:
		try:
			os.startfile(Path(self.excel_path).resolve())
		except OSError as error:
			messagebox.showerror("Could not open Excel", str(error))


def launch_gui(excel_path: str = DEFAULT_EXCEL_FILE) -> None:
	root = tk.Tk()
	RankingsWindow(root, excel_path)
	root.mainloop()


def build_parser() -> argparse.ArgumentParser:
	parser = argparse.ArgumentParser(description="Rank and compare fantasy football players.")
	subparsers = parser.add_subparsers(dest="command")

	rankings = subparsers.add_parser("rankings", help="Show global or positional rankings.")
	rankings.add_argument("--position", choices=POSITIONS, help="Only show one position.")
	rankings.add_argument("--sort", choices=("global", "position"), default="global")
	rankings.add_argument("--limit", type=int, default=25)
	rankings.add_argument("--excel", default=DEFAULT_EXCEL_FILE, help="Read rankings from an Excel workbook.")

	trade = subparsers.add_parser("trade", help="Compare the value of two sides of a trade.")
	trade.add_argument("--give", nargs="+", required=True, help="Player names you give away.")
	trade.add_argument("--get", nargs="+", required=True, help="Player names you receive.")
	trade.add_argument("--excel", default=DEFAULT_EXCEL_FILE, help="Read players from an Excel workbook.")

	fetch = subparsers.add_parser("fetch", help="Fetch current NFL players from Sleeper and export them to Excel.")
	fetch.add_argument("--output", default=DEFAULT_EXCEL_FILE, help="Excel file to create.")

	gui = subparsers.add_parser("gui", help="Open the rankings window.")
	gui.add_argument("--excel", default=DEFAULT_EXCEL_FILE, help="Read rankings from an Excel workbook.")
	return parser


def main() -> None:
	args = build_parser().parse_args()
	try:
		if args.command is None:
			launch_gui()
		elif args.command == "gui":
			launch_gui(args.excel)
		elif args.command == "rankings":
			if args.limit < 1:
				raise ValueError("--limit must be at least 1")
			print_rankings(load_players(args.excel), args.position, args.limit, args.sort)
		elif args.command == "trade":
			evaluate_trade(args.give, args.get, load_players(args.excel))
		else:
			players = fetch_api_players()
			write_excel(players, args.output)
			print(f"Saved {len(players)} players to {Path(args.output).resolve()}")
	except OSError as error:
		raise SystemExit(f"Error accessing data: {error}") from error
	except (KeyError, TypeError, ValueError) as error:
		raise SystemExit(f"Error reading player data: {error}") from error


if __name__ == "__main__":
	main()
